"""Generate the final report as a .docx (11pt) so it can be edited."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


def add_para(doc, text, bold=False, italic=False, size=11, align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 14, 2: 12}
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(sizes.get(level, 12))
    pf = p.paragraph_format
    pf.space_before = Pt(6 if level == 1 else 5)
    pf.space_after = Pt(3)
    return p


def add_table(doc, rows, col_widths_cm, header=True, font_size=10):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for i, w in enumerate(col_widths_cm):
        for cell in table.columns[i].cells:
            cell.width = Cm(w)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.size = Pt(font_size)
            if header and ri == 0:
                r.bold = True
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    # margin after table
    add_para(doc, "", size=4, space_after=2)
    return table


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--out", default="/home/elicer/project/report.docx")
    p.add_argument("--wer_json", default="/home/elicer/project/runs/fs2/eval/test_500/wer_summary.json")
    p.add_argument("--wer_jsonl", default="/home/elicer/project/runs/fs2/eval/test_500/wer_results.jsonl")
    p.add_argument("--train_curve", default="/home/elicer/project/report_assets/train_curves.png")
    p.add_argument("--loss_total", default="/home/elicer/project/report_assets/loss_total.png")
    args = p.parse_args()

    doc = Document()

    # tighten page margins
    for s in doc.sections:
        s.left_margin = Cm(1.7)
        s.right_margin = Cm(1.7)
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)

    # default style 11pt
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_heading(doc, "과제 3 — Duration 기반 TTS (FastSpeech2 + HiFi-GAN)", level=1)
    add_para(doc, "김효은 · 음성언어처리 · 2026", size=9.5, space_after=8)

    # 1
    add_heading(doc, "1. 개요", level=2)
    add_para(doc,
        "LJSpeech (훈련 12,500 / 검증 100 / 테스트 500)를 대상으로 FastSpeech2 방식의 "
        "duration 기반 TTS를 학습하였다. 음소 시퀀스와 음소별 duration은 제공된 "
        "IPA 압축 매니페스트에서 직접 읽어오므로 별도의 forced alignment는 수행하지 않았다. "
        "추론 시에는 ground-truth duration 레이블 대신 duration predictor의 예측값으로 "
        "프레임 수를 결정한다. 80-bin 멜 스펙트로그램은 사전 학습된 HiFi-GAN (LJ_FT_T2_V1)으로 "
        "보코딩하였다. 과제 허용 범위에 따라 duration predictor만 사용하고 "
        "pitch / energy predictor는 생략하였다.")

    # 2
    add_heading(doc, "2. 모델 구조", level=2)
    add_table(doc, [
        ["블록", "설정"],
        ["Embedding", "vocab 50 (IPA 46개 + pad/bos/eos/unk), d=256"],
        ["Encoder", "FFT 블록 4개: 2-head MHA + ConvFFN (kernel 9), d=256, FF 1024, dropout 0.1"],
        ["Duration predictor", "Conv1d (k=3, d=256) 2개 + LayerNorm + ReLU → Linear, dropout 0.5; 인코더 출력에 stop-gradient 적용"],
        ["Length regulator", "정수 duration만큼 인코더 hidden state를 반복 (torch.repeat_interleave)"],
        ["Decoder", "FFT 블록 4개 (인코더와 동일 설정)"],
        ["Mel head + postnet", "Linear 256 → 80, 이후 5층 1D-conv postnet (tanh + BN, 잔차 연결)"],
    ], col_widths_cm=[3.7, 13.6], font_size=10)
    add_para(doc,
        "인코더와 디코더 입력 모두에 사인파 위치 인코딩(sinusoidal positional encoding)을 더한다. "
        "총 학습 가능 파라미터 수: 41.49 M. 멜 스펙트로그램은 HiFi-GAN 설정과 동일하게 구성하였다 "
        "(22.05 kHz, n_fft 1024, hop 256, win 1024, 80 mels, f_min 0, f_max 8000, "
        "log-magnitude + dynamic-range compression).")

    # 3
    add_heading(doc, "3. 학습", level=2)
    add_para(doc,
        "음소 duration(초)은 round(end·sr/hop) − round(start·sr/hop) 방식으로 정수 멜 프레임으로 변환하여 "
        "음소 프레임 합이 멜 길이와 정확히 일치하도록 하였다. "
        "옵티마이저: AdamW (β = 0.9 / 0.98, ε = 1e-9, weight decay 1e-6). "
        "학습률 스케줄: Noam (d_model 256, warmup 4,000). 배치 크기 32, 그래디언트 클리핑 1.0. "
        "손실 함수는 postnet 이전 mel L1, postnet 이후 mel L1, 그리고 "
        "실제 음소 위치에서 예측 log-duration과 log(d_gt + 1) 간의 MSE를 합산한다. "
        "duration 손실은 인코더로 역전파되지 않도록 인코더 출력을 detach한 후 "
        "duration predictor에 입력한다. 총 50,000 스텝 학습 (~4.8시간, A100 MIG 2g.20gb).")

    add_heading(doc, "3.1 실험 추적 (wandb)", level=2)
    add_para(doc,
        "모든 학습 지표를 wandb에 기록하였다 (run ID: q5bpag6e). "
        "기록 항목: train/val 손실, 학습률, 그래디언트 norm. "
        "URL: https://wandb.ai/hyoeunkim-sungkyunkwan-university/slp-fs2-ljspeech/runs/q5bpag6e")

    # images
    if Path(args.train_curve).exists():
        doc.add_picture(args.train_curve, width=Cm(17.5))
    if Path(args.loss_total).exists():
        doc.add_picture(args.loss_total, width=Cm(11.0))

    add_para(doc, "스텝 50,000 최종 지표:", bold=True, space_after=2)
    add_table(doc, [
        ["지표", "훈련 (step 50,000)", "검증"],
        ["total loss",   "1.024", "1.577"],
        ["mel L1 (before)", "0.476", "0.750"],
        ["mel L1 (after)",  "0.465", "0.752"],
        ["duration MSE",   "0.083", "0.075"],
        ["lr",             "2.80 × 10⁻⁴", "—"],
        ["grad-norm",      "1.07",        "—"],
    ], col_widths_cm=[4.3, 5.5, 4.5], font_size=10)

    # 4
    add_heading(doc, "4. 평가 (Whisper-large-v3-turbo)", level=2)
    with open(args.wer_json) as f:
        wer = json.load(f)
    add_table(doc, [
        ["샘플 수", "WER", "대체(S)", "삽입(I)", "삭제(D)", "정답(H)"],
        [str(wer["n_samples"]), f"{wer['wer']*100:.2f}%",
         str(wer["substitutions"]), str(wer["insertions"]),
         str(wer["deletions"]), str(wer["hits"])],
    ], col_widths_cm=[2.0, 2.2, 2.0, 2.0, 2.0, 2.0], font_size=10)
    add_para(doc,
        "테스트셋 500개 발화 각각을 예측 duration으로 end-to-end 합성한 후 "
        "Whisper (large-v3-turbo, 영어, FP16)로 전사하였다. "
        "레퍼런스와 가설 텍스트는 모두 소문자로 변환하고 영숫자·아포스트로피 이외의 문자를 제거한 뒤 "
        "jiwer.process_words로 WER을 계산하였다.")

    # 5
    add_heading(doc, "5. 생성 샘플 (테스트 발화 5개)", level=2)
    add_para(doc,
        "22,050 Hz로 합성하였다. 레퍼런스 텍스트와 Whisper 전사 결과는 아래 표와 같으며, "
        "wav 파일은 제출물의 samples/ 폴더에 포함되어 있다.")
    sample_ids = ["LJ012-0091", "LJ014-0066", "LJ003-0011", "LJ002-0272", "LJ015-0205"]
    results = {}
    with open(args.wer_jsonl) as f:
        for line in f:
            r = json.loads(line)
            if r["id"] in sample_ids:
                results[r["id"]] = r
    rows = [["ID", "레퍼런스", "Whisper 전사"]]
    for uid in sample_ids:
        r = results.get(uid)
        if r is None:
            rows.append([uid, "(missing)", ""])
        else:
            rows.append([uid, r["ref"], r["hyp"]])
    add_table(doc, rows, col_widths_cm=[2.6, 7.4, 7.4], font_size=9)

    # 6
    add_heading(doc, "6. 고찰", level=2)
    add_para(doc,
        "• 검증 mel-L1은 0.75 부근에서 수렴하는 반면 훈련 손실은 계속 감소하여 "
        "소폭의 과적합이 관찰되었다. 그럼에도 합성 음성의 가청성은 유지되었다.")
    add_para(doc,
        "• 보코더로 LJ_FT_T2_V1을 선택한 이유는, 이 모델이 Tacotron2 생성 멜을 대상으로 "
        "파인튜닝되었기 때문이다. Tacotron2 출력 멜의 특성이 ground-truth 멜보다 "
        "FastSpeech2 출력과 유사하여 보코딩 품질이 더 좋았다.")
    add_para(doc,
        "• WER 오류 대부분은 음운론적 혼동이나 Whisper 정규화로 인한 것이며 "
        "(예: \"whitecross\" → \"white cross\", \"eighteen fifteen\" → \"1815\"), "
        "TTS 자체의 아티팩트로 보기 어렵다.")

    doc.save(args.out)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
